"""
NVSHOP.LK Comprehensive Academic Thesis Report Generator
Produces a high-standard, publication-quality academic dissertation (.docx)
covering every architectural, technical, functional, database, and implementation detail.
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell padding in twips (1/20th of a point)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m_name, m_val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m_name}')
        node.set(qn('w:w'), str(m_val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    """Set background color of a cell."""
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_table_borders(table, color="CBD5E1", sz="4"):
    """Apply elegant table borders."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="E2E8F0"/>'
        f'  <w:insideV w:val="none"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_header_footer(doc):
    """Add professional running header and footer."""
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("NVSHOP.LK – Honours Engineering Thesis Report | Second Year Software Project")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(100, 116, 139) # Slate 500
        
        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        frun1 = fp.add_run("Department of Software Engineering  •  Confidential & Academic Dissertation  •  ")
        frun1.font.name = "Calibri"
        frun1.font.size = Pt(8.5)
        frun1.font.color.rgb = RGBColor(100, 116, 139)
        frun2 = fp.add_run("Page Ref: NVSHOP-THESIS-2026")
        frun2.font.name = "Calibri"
        frun2.font.size = Pt(8.5)
        frun2.font.bold = True
        frun2.font.color.rgb = RGBColor(30, 58, 138)

def format_run(run, font_name="Calibri", size_pt=11, bold=False, italic=False, color_rgb=(51, 65, 85)):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)

def add_paragraph(doc, text, bold_prefix="", space_after=6, space_before=0, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = line_spacing
    
    if bold_prefix:
        r_prefix = p.add_run(bold_prefix)
        format_run(r_prefix, font_name="Calibri", size_pt=11, bold=True, color_rgb=(15, 23, 42))
    
    if text:
        r_text = p.add_run(text)
        format_run(r_text, font_name="Calibri", size_pt=11, color_rgb=(51, 65, 85))
    return p

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    format_run(r, font_name="Calibri Light", size_pt=18, bold=True, color_rgb=(30, 58, 138)) # Navy #1E3A8A
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    format_run(r, font_name="Calibri Light", size_pt=14, bold=True, color_rgb=(37, 99, 235)) # Royal Blue #2563EB
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    format_run(r, font_name="Calibri", size_pt=12, bold=True, color_rgb=(15, 23, 42)) # Dark Slate #0F172A
    return p

def add_bullet(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_prefix = p.add_run(bold_prefix)
        format_run(r_prefix, font_name="Calibri", size_pt=10.5, bold=True, color_rgb=(15, 23, 42))
    r_text = p.add_run(text)
    format_run(r_text, font_name="Calibri", size_pt=10.5, color_rgb=(51, 65, 85))
    return p

def add_callout(doc, title, text, icon="ℹ️"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_shading(cell, "F0F9FF") # Light Sky Blue
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    # Left thick border
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="0284C7"/>'
        f'  <w:top w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:bottom w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.space_before = Pt(0)
    r_title = p.add_run(f"{icon}  {title}\n")
    format_run(r_title, font_name="Calibri", size_pt=11, bold=True, color_rgb=(3, 105, 161))
    
    r_body = p.add_run(text)
    format_run(r_body, font_name="Calibri", size_pt=10, italic=False, color_rgb=(51, 65, 85))
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_styled_table(doc, headers, rows_data, col_widths=None):
    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    
    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_shading(hdr_cells[i], "1E3A8A") # Navy
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            format_run(r, font_name="Calibri", size_pt=10, bold=True, color_rgb=(255, 255, 255))
        if col_widths and i < len(col_widths):
            hdr_cells[i].width = Inches(col_widths[i])
            
    # Data Rows
    for r_idx, row_values in enumerate(rows_data):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            row_cells[c_idx].text = str(val)
            set_cell_shading(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=90, bottom=90, left=140, right=140)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                format_run(r, font_name="Calibri", size_pt=9.5, color_rgb=(51, 65, 85))
            if col_widths and c_idx < len(col_widths):
                row_cells[c_idx].width = Inches(col_widths[c_idx])
                
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table

print("Generator helper library loaded.")
